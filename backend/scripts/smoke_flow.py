"""Full workflow with simulated cloud HTTP; no real API credentials are tested."""
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
import httpx
from openai import AsyncOpenAI, OpenAI
from PIL import Image
import pymupdf

from app import main
from app.config import Settings
from app.services.session import _attempts, _sessions
from app.services.state import SQLiteMap


def make_docx() -> bytes:
    document = Document()
    document.add_heading("脱敏候选人简历", 0)
    document.add_paragraph("使用 Python 和 SQL 完成数据分析项目，处理 10 万条数据。")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def run() -> None:
    requests: list[dict] = []

    def cloud_response(request: httpx.Request) -> httpx.Response:
        assert request.url.host in {"dashscope.aliyuncs.com", "dashscope-intl.aliyuncs.com"}, request.url
        assert request.url.path == "/compatible-mode/v1/chat/completions", request.url
        body = json.loads(request.content)
        token = request.headers["authorization"].removeprefix("Bearer ")
        requests.append({"host": request.url.host, "key": token, "body": body})
        if token == "sk-invalid-secret":
            return httpx.Response(401, json={"error": {"message": f"invalid credential: {token}", "type": "invalid_api_key"}})
        system = body["messages"][0]["content"]
        if body.get("max_tokens") == 1:
            result = "OK"
        elif "OCR" in system:
            result = "PYTHON SQL EXCEL DATA ANALYST"
        elif "简历结构化" in system:
            result = {"education": [{"degree": "本科"}], "skills": [{"name": "Python", "normalized_name": "python"}, {"name": "SQL", "normalized_name": "sql"}], "experiences": [], "projects": [{"name": "数据分析项目"}], "preferences": {"cities": ["上海"]}, "warnings": []}
        elif "简历辅导员" in system:
            result = {"optional_improve": ["补充项目结果"], "checklist": ["核对事实"]}
        elif "模拟面试官" in system:
            result = {"questions": [{"question": question, "keywords": []} for question in ["请介绍一个项目及其结果。", "请说明你的数据分析方法。", "请说明你如何验证项目结果。", "请说明一次具体的团队协作。"]]}
        elif "练习反馈" in system:
            result = {"score": 4.0, "feedback": "回答包含项目证据。"}
        else:
            raise AssertionError(f"Unexpected cloud request: {system}")
        content = result if isinstance(result, str) else json.dumps(result, ensure_ascii=False)
        return httpx.Response(200, json={"id": "simulated", "object": "chat.completion", "created": 0, "model": body["model"], "choices": [{"index": 0, "message": {"role": "assistant", "content": content}, "finish_reason": "stop"}]})

    def async_client(**kwargs):
        return AsyncOpenAI(**kwargs, http_client=httpx.AsyncClient(transport=httpx.MockTransport(cloud_response)))

    def sync_client(**kwargs):
        return OpenAI(**kwargs, http_client=httpx.Client(transport=httpx.MockTransport(cloud_response)))

    previous_storage, previous_state, previous_key = main.storage, main.settings.state_db_path, main.settings.qwen_api_key
    main.settings.qwen_api_key = ""
    _attempts.clear()
    _sessions.clear()
    with TemporaryDirectory(prefix="careerpilot-cloud-smoke-") as temporary:
        root = Path(temporary)
        state_path = root / "state.sqlite3"
        SQLiteMap(state_path, "persistence-check")["profile-1"] = {"status": "confirmed"}
        assert SQLiteMap(state_path, "persistence-check")["profile-1"]["status"] == "confirmed"
        main.storage = main.LocalStorage(root / "uploads")
        main.configure_state(state_path)
        try:
            with patch("app.services.llm.AsyncOpenAI", async_client), patch("app.services.llm.OpenAI", sync_client), TestClient(main.app) as client, TestClient(main.app) as stranger:
                assert client.get("/api/v1/model-config").json()["chat_configured"] is False
                assert client.post("/api/v1/model/ping").status_code == 503
                oversized = client.post("/api/v1/model/configure", content=json.dumps({"api_key": "x" * 5000}), headers={"Content-Type": "application/json"})
                assert oversized.status_code == 413 and "x" * 100 not in oversized.text, oversized.text
                invalid = client.post("/api/v1/model/configure", json={"api_key": "sk-invalid-secret"})
                assert invalid.status_code == 502 and "sk-invalid-secret" not in invalid.text, invalid.text
                assert not client.get("/api/v1/model-config").json()["chat_configured"]
                for payload in [{"api_key": "sk invalid-secret"}, {"api_key": "sk-browser-one", "base_url": "http://localhost:1234/v1"}, {"api_key": "sk-browser-one", "region": "us"}]:
                    denied = client.post("/api/v1/model/configure", json=payload)
                    assert denied.status_code == 422 and payload["api_key"] not in denied.text, denied.text
                configured = client.post("/api/v1/model/configure", json={"api_key": "sk-browser-one"})
                assert configured.status_code == 200 and configured.json()["chat_configured"], configured.text
                assert configured.json()["region"] == "china" and "sk-browser-one" not in configured.text
                assert configured.headers["cache-control"] == "no-store"
                assert "careerpilot_session=" in client.get("/api/v1/model-config").request.headers["cookie"]
                assert not stranger.get("/api/v1/model-config").json()["chat_configured"]
                second = stranger.post("/api/v1/model/configure", json={"api_key": "sk-browser-two", "region": "singapore", "model": "qwen-turbo"})
                assert second.status_code == 200, second.text
                assert stranger.post("/api/v1/model/ping").json()["model"] == "qwen-turbo"
                assert requests[-1]["key"] == "sk-browser-two" and requests[-1]["host"] == "dashscope-intl.aliyuncs.com"
                assert client.post("/api/v1/model/ping").json()["model"] == "qwen-plus"
                assert requests[-1]["key"] == "sk-browser-one"
                failed_change = client.post("/api/v1/model/configure", json={"api_key": "sk-invalid-secret", "model": "qwen-turbo"})
                assert failed_change.status_code == 502 and client.get("/api/v1/model-config").json()["model"] == "qwen-plus"
                assert client.post("/api/v1/resumes/upload", files={"file": ("resume.docx", make_docx())}).status_code == 400
                upload = client.post("/api/v1/resumes/upload", files={"file": ("resume.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, data={"privacy_acknowledged": "true"})
                assert upload.status_code == 201, upload.text
                resume_id = upload.json()["resume_id"]
                assert stranger.get(f"/api/v1/resumes/{resume_id}/blocks").status_code == 404
                draft = client.post(f"/api/v1/resumes/{resume_id}/profile")
                assert draft.status_code == 200, draft.text
                confirmed = client.post(f"/api/v1/resumes/{resume_id}/profile/confirm", json=draft.json())
                assert confirmed.status_code == 200 and confirmed.json()["status"] == "confirmed", confirmed.text
                jobset = client.post("/api/v1/jobsets", json={"name": "烟测岗位"})
                assert jobset.status_code == 201, jobset.text
                jobset_id = jobset.json()["jobset_id"]
                assert stranger.get(f"/api/v1/jobsets/{jobset_id}").status_code == 404
                imported = client.post(f"/api/v1/jobsets/{jobset_id}/jobs/import", json={"jobs": [{"title": "数据分析师", "company": "示例公司", "location": "上海", "education_required": "本科", "required_skills": ["Python", "SQL"]}]})
                assert imported.status_code == 200, imported.text
                matches = client.post("/api/v1/matches/run", json={"resume_id": resume_id, "jobset_id": jobset_id})
                assert matches.status_code == 200 and matches.json()["results"][0]["eligible"], matches.text
                assert "semantic_similarity" not in matches.json()["results"][0]["score_breakdown"]
                target = {"resume_id": resume_id, "jobset_id": jobset_id, "job_title": "数据分析师"}
                diagnosis = client.post("/api/v1/reports/diagnosis", json=target)
                assert diagnosis.status_code == 200 and diagnosis.json()["engine"] == "model", diagnosis.text
                interview = client.post("/api/v1/interviews", json=target)
                assert interview.status_code == 201 and interview.json()["engine"] == "model", interview.text
                session = interview.json()
                assert stranger.get(f"/api/v1/interviews/{session['session_id']}").status_code == 404
                while session["status"] == "active":
                    answer = client.post(f"/api/v1/interviews/{session['session_id']}/answer", json={"answer": "我使用 Python 与 SQL 处理 10 万条数据，最终提升报表效率 30%。"})
                    assert answer.status_code == 200, answer.text
                    session = answer.json()
                assert session["overall_score"] == 4.0
                export = client.post("/api/v1/exports", json={"type": "diagnosis", "resource_id": diagnosis.json()["report_id"], "format": "json"})
                assert export.status_code == 201 and "数据分析师" in export.json()["content"], export.text
                assert stranger.get(f"/api/v1/exports/{export.json()['export_id']}").status_code == 404
                interview_export = client.post("/api/v1/exports", json={"type": "interview", "resource_id": session["session_id"], "format": "html"})
                assert interview_export.status_code == 201 and "模拟面试反馈" in interview_export.json()["content"], interview_export.text
                image_buffer = BytesIO()
                Image.new("RGB", (800, 160), "white").save(image_buffer, format="PNG")
                image_bytes = image_buffer.getvalue()
                image_upload = client.post("/api/v1/resumes/upload", files={"file": ("ocr.png", image_bytes, "image/png")}, data={"privacy_acknowledged": "true"})
                assert image_upload.status_code == 201, image_upload.text
                assert requests[-1]["key"] == "sk-browser-one" and requests[-1]["body"]["model"] == "qwen-vl-plus"
                image_blocks = client.get(f"/api/v1/resumes/{image_upload.json()['resume_id']}/blocks").json()["blocks"]
                assert "PYTHON" in " ".join(block["text"] for block in image_blocks)
                with pymupdf.open() as document:
                    page = document.new_page(width=800, height=160)
                    page.insert_image(pymupdf.Rect(0, 0, 800, 160), stream=image_bytes)
                    pdf_bytes = document.tobytes()
                pdf_upload = client.post("/api/v1/resumes/upload", files={"file": ("scanned.pdf", pdf_bytes, "application/pdf")}, data={"privacy_acknowledged": "true"})
                assert pdf_upload.status_code == 201, pdf_upload.text
                blocks = client.get(f"/api/v1/resumes/{pdf_upload.json()['resume_id']}/blocks").json()["blocks"]
                assert "EXCEL" in " ".join(block["text"] for block in blocks)
                assert not stranger.delete("/api/v1/model/configure").json()["chat_configured"]
                assert client.get("/api/v1/model-config").json()["chat_configured"]
                demo = stranger.post("/api/v1/demo/china/start")
                assert demo.status_code == 201 and "合成" in demo.json()["notice"], demo.text
                assert all(job["synthetic"] for job in demo.json()["jobset"]["jobs"])
                assert not client.delete("/api/v1/model/configure").json()["chat_configured"]
                assert client.post("/api/v1/model/ping").status_code == 503
                limited = [client.post("/api/v1/model/configure", json={"api_key": "sk-invalid-secret"}).status_code for _ in range(8)]
                assert 429 in limited, limited
                for path in root.rglob("*"):
                    if path.is_file():
                        content = path.read_bytes()
                        assert b"sk-browser-one" not in content and b"sk-browser-two" not in content, path
                assert Settings(qwen_api_key="sk-test-key").qwen_base_url == "https://dashscope.aliyuncs.com/compatible-mode/v1"
                with patch.dict("os.environ", {"DASHSCOPE_API_KEY": "sk-server-default", "QWEN_API_KEY": "sk-alternative-default"}):
                    assert Settings(_env_file=None).qwen_api_key == "sk-server-default"
                with patch.dict("os.environ", {"DASHSCOPE_API_KEY": "", "QWEN_API_KEY": "sk-alternative-default"}):
                    assert Settings(_env_file=None).qwen_api_key == "sk-alternative-default"
                main.settings.qwen_api_key = "sk-server-default"
                with TestClient(main.app) as default_client:
                    assert default_client.get("/api/v1/model-config").json()["chat_configured"]
                    assert not default_client.delete("/api/v1/model/configure").json()["chat_configured"]
        finally:
            main.storage = previous_storage
            main.configure_state(previous_state)
            main.settings.qwen_api_key = previous_key
            _sessions.clear()
            _attempts.clear()
    print(f"Cloud HTTP simulation + browser isolation + complete career workflow: PASS ({len(requests)} simulated requests; no real cloud call)")


if __name__ == "__main__":
    run()
