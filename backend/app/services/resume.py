from collections.abc import Callable
from pathlib import Path

import pymupdf
from docx import Document


def initial_resume_blocks(
    path: Path,
    content_type: str,
    qwen_vision_ocr: Callable[[bytes, str], str] | None = None,
) -> list[dict[str, str | int]]:
    """Parse document text; image-only pages use Alibaba Cloud Qwen-VL."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        blocks: list[dict[str, str | int]] = []
        with pymupdf.open(path) as document:
            for page_number, page in enumerate(document[:20], start=1):
                page_text = [block[4].strip() for block in page.get_text("blocks") if block[4].strip()]
                if page_text:
                    blocks.extend({"block_id": f"page_{page_number}_block_{index}", "text": text[:4000], "page": page_number} for index, text in enumerate(page_text, start=1))
                    continue
                if qwen_vision_ocr is None:
                    raise RuntimeError("请先连接千问云 API，再处理扫描 PDF")
                text = qwen_vision_ocr(page.get_pixmap(dpi=200, alpha=False).tobytes("png"), "image/png").strip()
                if text:
                    blocks.append({"block_id": f"page_{page_number}_vision_ocr", "text": text[:12000], "page": page_number})
        return blocks or [{"block_id": "ocr_empty", "text": "未能识别到简历文字，请检查文件后重试。", "page": 1}]

    if suffix == ".docx":
        document = Document(path)
        blocks: list[dict[str, str | int]] = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if text:
                blocks.append({"block_id": f"paragraph_{index}", "text": text[:4000], "page": 1})
        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    blocks.append({"block_id": f"table_{table_index}_row_{row_index}", "text": text[:4000], "page": 1})
        return blocks

    if content_type.startswith("text/"):
        text = path.read_text(encoding="utf-8", errors="replace")
        return [{"block_id": "text_1", "text": text[:20000], "page": 1}]

    if qwen_vision_ocr is None:
        raise RuntimeError("请先连接千问云 API，再处理图片简历")
    media_type = "image/png" if suffix == ".png" else "image/jpeg"
    text = qwen_vision_ocr(path.read_bytes(), media_type).strip()
    return [{"block_id": "vision_ocr_1", "text": text[:12000], "page": 1}] if text else [
        {"block_id": "ocr_empty", "text": "未能识别到简历文字，请检查图片后重试。", "page": 1}
    ]

