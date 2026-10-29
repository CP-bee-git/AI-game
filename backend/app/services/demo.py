"""A source-labelled Chinese demo; no model call or real person data."""
import json
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


DEMO_FILE = Path(__file__).resolve().parents[1] / "data" / "china_demo.json"


def load_china_demo() -> dict:
    return json.loads(DEMO_FILE.read_text(encoding="utf-8"))


def make_demo_resume(data: dict) -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.element.get_or_add_rPr().append(fonts)
    for index, paragraph in enumerate(data["resume_paragraphs"]):
        if index == 0:
            document.add_heading(paragraph, 0)
        else:
            document.add_paragraph(paragraph)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
